import os
import sqlite3
import tempfile
import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import excel_import


class ExcelImportTests(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix='.db')
        os.close(fd)
        conn = sqlite3.connect(self.db_path)
        conn.execute('PRAGMA foreign_keys = ON')
        conn.executescript('''
            CREATE TABLE products (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL);
            CREATE TABLE modules (id INTEGER PRIMARY KEY AUTOINCREMENT, product_id INTEGER NOT NULL, name TEXT NOT NULL, UNIQUE(product_id, name));
            CREATE TABLE test_cases (id INTEGER PRIMARY KEY AUTOINCREMENT, module_id INTEGER NOT NULL, case_title TEXT NOT NULL, preconditions TEXT, steps TEXT, expected_result TEXT, remark TEXT, priority TEXT, updated_at TIMESTAMP);
        ''')
        conn.close()
        self.connection_patch = patch('excel_import.get_connection', self._connection)
        self.connection_patch.start()

    def tearDown(self):
        self.connection_patch.stop()
        os.remove(self.db_path)

    def _connection(self):
        conn = sqlite3.connect(self.db_path)
        conn.execute('PRAGMA foreign_keys = ON')
        return conn

    def _workbook_bytes(self, rows):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(excel_import.IMPORT_HEADERS)
        for row in rows:
            worksheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_template_has_expected_headers(self):
        workbook = load_workbook(excel_import.build_template())
        self.assertEqual([cell.value for cell in workbook.active[1]], excel_import.IMPORT_HEADERS)

    def test_preview_and_confirm_import_preserve_newlines(self):
        payload = self._workbook_bytes([[
            'Web', 'Login', '登入成功', '已註冊', '1. 輸入帳號\n2. 點擊登入',
            '進入首頁', '', '備註\n第二行',
        ]])
        preview = excel_import.validate_import(payload, 'cases.xlsx')
        self.assertTrue(preview['valid'])
        self.assertEqual(preview['new_products'], ['Web'])
        with self.assertRaises(excel_import.ImportValidationError):
            excel_import.import_file(payload, 'cases.xlsx')
        result = excel_import.import_file(payload, 'cases.xlsx', confirm_new=True)
        self.assertEqual(result['created_testcases'], 1)
        conn = self._connection()
        stored = conn.execute('SELECT steps, remark, priority FROM test_cases').fetchone()
        conn.close()
        self.assertEqual(stored, ('1. 輸入帳號\n2. 點擊登入', '備註\n第二行', 'Medium'))

    def test_invalid_headers_and_duplicate_cases_are_rejected(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(['Wrong Header'])
        buffer = BytesIO()
        workbook.save(buffer)
        with self.assertRaises(excel_import.ImportValidationError):
            excel_import.validate_import(buffer.getvalue(), 'cases.xlsx')

        payload = self._workbook_bytes([
            ['Web', 'Login', '登入成功', '', '', '', 'High', ''],
            ['Web', 'Login', '登入成功', '', '', '', 'High', ''],
        ])
        preview = excel_import.validate_import(payload, 'cases.xlsx')
        self.assertFalse(preview['valid'])
        self.assertIn('不可重複', preview['errors'][0]['message'])

    def test_existing_case_is_rejected_but_existing_product_and_module_are_reused(self):
        conn = self._connection()
        conn.execute("INSERT INTO products (name) VALUES ('Web')")
        product_id = conn.execute('SELECT id FROM products WHERE name = ?', ('Web',)).fetchone()[0]
        conn.execute("INSERT INTO modules (product_id, name) VALUES (?, 'Login')", (product_id,))
        module_id = conn.execute('SELECT id FROM modules WHERE product_id = ?', (product_id,)).fetchone()[0]
        conn.execute("INSERT INTO test_cases (module_id, case_title) VALUES (?, '既有 Case')", (module_id,))
        conn.commit()
        conn.close()

        duplicate = self._workbook_bytes([['Web', 'Login', '既有 Case', '', '', '', '', '']])
        duplicate_preview = excel_import.validate_import(duplicate, 'cases.xlsx')
        self.assertFalse(duplicate_preview['valid'])
        self.assertIn('系統已有', duplicate_preview['errors'][0]['message'])

        new_case = self._workbook_bytes([['Web', 'Login', '新的 Case', '', '', '', '', '']])
        preview = excel_import.validate_import(new_case, 'cases.xlsx')
        self.assertTrue(preview['valid'])
        self.assertEqual(preview['new_products'], [])
        self.assertEqual(preview['new_modules'], [])
        result = excel_import.import_file(new_case, 'cases.xlsx')
        self.assertEqual(result['created_products'], 0)
        self.assertEqual(result['created_modules'], 0)


if __name__ == '__main__':
    unittest.main()
