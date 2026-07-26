from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from database import get_connection


IMPORT_HEADERS = [
    'Product', 'Module', 'Case', 'Preconditions',
    'Steps', 'Expected Result', 'Priority', 'Remark',
]
VALID_PRIORITIES = {'Low', 'Medium', 'High', 'Critical'}


class ImportValidationError(ValueError):
    def __init__(self, message, errors=None):
        super().__init__(message)
        self.errors = errors or []


def build_template():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'TestCases'
    worksheet.append(IMPORT_HEADERS)
    worksheet.append([
        'Web 1.0', 'Login', '使用有效帳密登入', '使用者已註冊',
        '1. 開啟登入頁\n2. 輸入帳號與密碼\n3. 點擊登入',
        '成功進入系統首頁', 'High', '可使用測試帳號驗證',
    ])
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    worksheet.freeze_panes = 'A2'
    worksheet.column_dimensions['A'].width = 20
    worksheet.column_dimensions['B'].width = 20
    worksheet.column_dimensions['C'].width = 30
    worksheet.column_dimensions['D'].width = 28
    worksheet.column_dimensions['E'].width = 42
    worksheet.column_dimensions['F'].width = 28
    worksheet.column_dimensions['G'].width = 14
    worksheet.column_dimensions['H'].width = 28
    worksheet.row_dimensions[2].height = 58
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _cell_text(value):
    if value is None:
        return ''
    return str(value)


def _existing_data(conn):
    products = {}
    modules = {}
    cases = set()
    for product_id, product_name in conn.execute('SELECT id, name FROM products ORDER BY id ASC'):
        products.setdefault((product_name or '').strip(), product_id)
    for module_id, product_id, module_name in conn.execute(
        'SELECT id, product_id, name FROM modules ORDER BY id ASC'
    ):
        modules.setdefault((product_id, (module_name or '').strip()), module_id)
    for product_name, module_name, case_title in conn.execute(
        '''SELECT p.name, m.name, t.case_title
           FROM test_cases t
           JOIN modules m ON m.id = t.module_id
           JOIN products p ON p.id = m.product_id'''
    ):
        cases.add(((product_name or '').strip(), (module_name or '').strip(), (case_title or '').strip()))
    return products, modules, cases


def parse_import_file(file_bytes, filename):
    if not filename or not filename.lower().endswith('.xlsx'):
        raise ImportValidationError('僅支援 .xlsx 格式的 Excel 檔案。')
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
        worksheet = workbook.worksheets[0]
    except Exception as error:
        raise ImportValidationError('Excel 檔案無法解析，請確認檔案未損毀且為有效的 .xlsx 格式。') from error

    headers = [_cell_text(cell.value).strip() for cell in next(worksheet.iter_rows(min_row=1, max_row=1), [])]
    if headers != IMPORT_HEADERS:
        raise ImportValidationError(
            'Excel 第一列欄位必須依序為：' + '、'.join(IMPORT_HEADERS) + '。'
        )

    rows = []
    errors = []
    for excel_row, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(values[:len(IMPORT_HEADERS)])
        if len(values) < len(IMPORT_HEADERS):
            values.extend([None] * (len(IMPORT_HEADERS) - len(values)))
        text_values = [_cell_text(value) for value in values]
        if not any(value.strip() for value in text_values):
            continue
        product, module, case_title = (text_values[index].strip() for index in range(3))
        preconditions, steps, expected_result = text_values[3:6]
        priority = text_values[6].strip() or 'Medium'
        remark = text_values[7]
        missing = [label for label, value in [('Product', product), ('Module', module), ('Case', case_title)] if not value]
        if missing:
            errors.append({'row': excel_row, 'message': '必填欄位不可空白：' + '、'.join(missing)})
        if priority not in VALID_PRIORITIES:
            errors.append({'row': excel_row, 'message': 'Priority 僅可為 Low、Medium、High、Critical，或留白。'})
        rows.append({
            'row': excel_row, 'product': product, 'module': module, 'case': case_title,
            'preconditions': preconditions, 'steps': steps,
            'expected_result': expected_result, 'priority': priority, 'remark': remark,
        })
    if not rows:
        errors.append({'row': 2, 'message': 'Excel 至少需要一筆 TestCase 資料。'})
    return rows, errors


def validate_import(file_bytes, filename):
    rows, errors = parse_import_file(file_bytes, filename)
    seen_cases = set()
    for row in rows:
        case_key = (row['product'], row['module'], row['case'])
        if case_key in seen_cases:
            errors.append({'row': row['row'], 'message': '同一 Product、Module 下的 Case 不可重複。'})
        seen_cases.add(case_key)

    conn = get_connection()
    try:
        existing_products, existing_modules, existing_cases = _existing_data(conn)
    finally:
        conn.close()
    for row in rows:
        if (row['product'], row['module'], row['case']) in existing_cases:
            errors.append({'row': row['row'], 'message': '系統已有相同 Product、Module 與 Case 的 TestCase。'})

    new_products = []
    new_modules = []
    planned_product_ids = dict(existing_products)
    planned_module_keys = {
        (product_name, module_name)
        for product_id, module_name in existing_modules
        for product_name, existing_id in existing_products.items()
        if existing_id == product_id
    }
    for row in rows:
        product_name = row['product']
        module_key = (product_name, row['module'])
        if product_name not in planned_product_ids:
            planned_product_ids[product_name] = None
            new_products.append(product_name)
        if module_key not in planned_module_keys:
            planned_module_keys.add(module_key)
            new_modules.append({'product': product_name, 'module': row['module']})

    return {
        'rows': rows,
        'errors': errors,
        'new_products': new_products,
        'new_modules': new_modules,
        'valid': not errors,
    }


def import_file(file_bytes, filename, confirm_new=False):
    preview = validate_import(file_bytes, filename)
    if preview['errors']:
        raise ImportValidationError('Excel 資料驗證失敗。', preview['errors'])
    if (preview['new_products'] or preview['new_modules']) and not confirm_new:
        raise ImportValidationError('此檔案需要建立新的 Product 或 Module，請先確認。')

    conn = get_connection()
    try:
        existing_products, existing_modules, existing_cases = _existing_data(conn)
        created_products = 0
        created_modules = 0
        created_cases = 0
        product_ids = dict(existing_products)
        module_ids = dict(existing_modules)
        for row in preview['rows']:
            product_name = row['product']
            product_id = product_ids.get(product_name)
            if product_id is None:
                cursor = conn.execute('INSERT INTO products (name) VALUES (?)', (product_name,))
                product_id = cursor.lastrowid
                product_ids[product_name] = product_id
                created_products += 1
            module_id = module_ids.get((product_id, row['module']))
            if module_id is None:
                cursor = conn.execute(
                    'INSERT INTO modules (product_id, name) VALUES (?, ?)',
                    (product_id, row['module']),
                )
                module_id = cursor.lastrowid
                module_ids[(product_id, row['module'])] = module_id
                created_modules += 1
            case_key = (product_name, row['module'], row['case'])
            if case_key in existing_cases:
                raise ImportValidationError('匯入期間發現重複 TestCase，請重新預覽檔案。')
            conn.execute(
                '''INSERT INTO test_cases (
                    module_id, case_title, preconditions, steps, expected_result, remark, priority, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)''',
                (module_id, row['case'], row['preconditions'], row['steps'],
                 row['expected_result'], row['remark'], row['priority']),
            )
            existing_cases.add(case_key)
            created_cases += 1
        conn.commit()
        return {
            'created_products': created_products,
            'created_modules': created_modules,
            'created_testcases': created_cases,
        }
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
